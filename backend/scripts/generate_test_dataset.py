"""Phase 0 Prompt 4 (option A) — deterministic synthetic test-dataset generator.

100% synthetic, zero real personal data. No network calls, no LLM calls:
pure-Python templates + seeded RNG only (free-tier safe, reproducible).

Usage (from repo root):
    python backend/scripts/generate_test_dataset.py --out test-data --resumes 50 --jds 20 --seed 42

Outputs in <out>/:
    resume_*.pdf / resume_*.docx   synthetic resumes
    jd_*.txt                       synthetic job descriptions
    dataset_manifest.json          counts + provenance metadata
    README_FOR_BUCKET.md           provenance note for the private Supabase bucket
"""

from __future__ import annotations

import argparse
import datetime
import json
import random
import textwrap
from pathlib import Path

# ---------------------------------------------------------------------------
# Role-family content (curated skill lists, titles, tasks per family)
# ---------------------------------------------------------------------------

ROLE_FAMILIES: list[str] = [
    "frontend",
    "backend",
    "data-analyst",
    "ml-engineer",
    "uiux",
    "product-manager",
    "marketing",
    "hr",
    "finance",
    "qa",
]

ROLE_DISPLAY = {
    "frontend": "Frontend Developer",
    "backend": "Backend Developer",
    "data-analyst": "Data Analyst",
    "ml-engineer": "ML Engineer",
    "uiux": "UI/UX Designer",
    "product-manager": "Product Manager",
    "marketing": "Marketing Specialist",
    "hr": "HR Generalist",
    "finance": "Finance Analyst",
    "qa": "QA Engineer",
}

ROLE_SKILLS: dict[str, list[str]] = {
    "frontend": ["HTML", "CSS", "JavaScript", "TypeScript", "React", "Next.js",
                 "Tailwind CSS", "Redux", "Vite", "Jest", "Accessibility (a11y)",
                 "Responsive Design", "Figma", "REST APIs"],
    "backend": ["Python", "FastAPI", "Node.js", "PostgreSQL", "Redis", "Docker",
                "REST APIs", "SQLAlchemy", "Authentication (JWT)", "Unit Testing",
                "CI/CD", "Message Queues", "Linux", "API Design"],
    "data-analyst": ["SQL", "Python", "pandas", "Excel", "Tableau", "Power BI",
                     "Data Cleaning", "Statistics", "A/B Testing", "Dashboarding",
                     "ETL Basics", "Data Visualization", "Google Sheets", "Reporting"],
    "ml-engineer": ["Python", "scikit-learn", "PyTorch", "TensorFlow", "pandas",
                    "NumPy", "Model Evaluation", "Feature Engineering", "MLOps Basics",
                    "Data Pipelines", "Experiment Tracking", "Statistics", "SQL", "Docker"],
    "uiux": ["Figma", "Wireframing", "Prototyping", "User Research", "Usability Testing",
             "Design Systems", "Interaction Design", "Information Architecture",
             "Adobe XD", "Accessibility (a11y)", "User Flows", "Visual Design",
             "FigJam", "Responsive Design"],
    "product-manager": ["Roadmapping", "User Stories", "Agile/Scrum", "Stakeholder Management",
                        "A/B Testing", "Analytics (Mixpanel)", "Prioritization (RICE)",
                        "Wireframing Basics", "SQL Basics", "Go-to-Market", "Customer Interviews",
                        "KPIs & Metrics", "Jira", "Prototyping"],
    "marketing": ["SEO", "Content Marketing", "Email Marketing", "Google Analytics",
                  "Social Media Marketing", "Copywriting", "Marketing Automation",
                  "A/B Testing", "Campaign Management", "CRM (HubSpot)", "Paid Ads Basics",
                  "Brand Guidelines", "Canva", "Reporting"],
    "hr": ["Recruiting", "Onboarding", "Employee Relations", "HR Policies", "Sourcing",
           "Interviewing", "Performance Reviews", "HRIS Basics", "Labor Law Basics",
           "Engagement Surveys", "Payroll Coordination", "Training & Development",
           "ATS Tools", "Excel"],
    "finance": ["Financial Modeling", "Excel", "Budgeting", "Forecasting", "Variance Analysis",
                "Accounting Basics", "SQL Basics", "Power BI", "Valuation Basics",
                "Cash Flow Analysis", "Reporting", "SAP Basics", "Statistics", "Auditing Basics"],
    "qa": ["Manual Testing", "Test Case Design", "Selenium", "API Testing (Postman)",
           "Bug Tracking (Jira)", "Regression Testing", "SQL Basics", "Test Plans",
           "Automation Basics (Pytest)", "Performance Testing Basics", "Agile/Scrum",
           "CI/CD Basics", "Linux Basics", "Documentation"],
}

ROLE_TITLES = {
    "frontend": ("Junior Frontend Developer", "Senior Frontend Developer"),
    "backend": ("Junior Backend Developer", "Senior Backend Developer"),
    "data-analyst": ("Junior Data Analyst", "Senior Data Analyst"),
    "ml-engineer": ("Junior ML Engineer", "Senior ML Engineer"),
    "uiux": ("Junior UI/UX Designer", "Senior UI/UX Designer"),
    "product-manager": ("Associate Product Manager", "Senior Product Manager"),
    "marketing": ("Marketing Associate", "Senior Marketing Specialist"),
    "hr": ("HR Associate", "Senior HR Generalist"),
    "finance": ("Junior Finance Analyst", "Senior Finance Analyst"),
    "qa": ("Junior QA Engineer", "Senior QA Engineer"),
}

ROLE_TASKS: dict[str, list[str]] = {
    "frontend": ["built responsive pages", "implemented reusable UI components",
                 "integrated REST APIs", "fixed cross-browser layout bugs",
                 "improved page accessibility", "migrated class components to hooks"],
    "backend": ["built REST API endpoints", "designed database schemas",
                "added background jobs", "fixed production API bugs",
                "implemented JWT authentication", "wrote unit tests for services"],
    "data-analyst": ["cleaned raw datasets", "built KPI dashboards",
                     "ran A/B test analyses", "automated weekly reports",
                     "wrote complex SQL queries", "presented insights to stakeholders"],
    "ml-engineer": ["trained classification models", "built feature pipelines",
                    "ran hyperparameter experiments", "evaluated model performance",
                    "cleaned training datasets", "deployed a demo inference API"],
    "uiux": ["created wireframes", "ran usability tests", "built clickable prototypes",
             "redesigned onboarding flows", "maintained a design system",
             "mapped user journeys"],
    "product-manager": ["wrote user stories", "prioritized the backlog",
                        "ran sprint planning", "interviewed users",
                        "defined success metrics", "coordinated release launches"],
    "marketing": ["wrote blog content", "ran email campaigns", "optimized landing pages",
                  "managed social media calendar", "tracked campaign analytics",
                  "coordinated webinar launches"],
    "hr": ["screened candidate resumes", "coordinated interviews", "ran onboarding sessions",
           "maintained HR records", "organized engagement activities",
           "drafted HR policy updates"],
    "finance": ["built budget models", "prepared monthly variance reports",
                "forecasted cash flows", "reconciled ledger entries",
                "automated Excel reports", "supported audit documentation"],
    "qa": ["wrote test cases", "executed regression suites", "logged defects in Jira",
           "tested REST APIs with Postman", "automated smoke tests",
           "verified bug fixes"],
}

ROLE_CERTS: dict[str, list[str]] = {
    "frontend": ["Demo Certified Frontend Associate", "Sample Web Accessibility Basics Certificate"],
    "backend": ["Demo Certified API Developer", "Sample Cloud Basics Certificate"],
    "data-analyst": ["Demo Certified Data Analyst", "Sample BI Tools Certificate"],
    "ml-engineer": ["Demo Certified ML Practitioner", "Sample Python for Data Certificate"],
    "uiux": ["Demo Certified UX Designer", "Sample Design Systems Certificate"],
    "product-manager": ["Demo Certified Scrum Associate", "Sample Product Analytics Certificate"],
    "marketing": ["Demo Certified Digital Marketer", "Sample SEO Fundamentals Certificate"],
    "hr": ["Demo Certified HR Associate", "Sample Recruiting Basics Certificate"],
    "finance": ["Demo Certified Finance Analyst", "Sample Excel Modeling Certificate"],
    "qa": ["Demo Certified QA Tester", "Sample Test Automation Basics Certificate"],
}

# Adjacent-role skills (sprinkled into JD nice-to-haves so matching is non-trivial).
ADJACENT_SKILLS = {
    "frontend": ["Figma", "Unit Testing", "CI/CD"],
    "backend": ["Docker", "React", "SQL"],
    "data-analyst": ["Tableau", "A/B Testing", "Python"],
    "ml-engineer": ["Docker", "SQL", "Statistics"],
    "uiux": ["Responsive Design", "Accessibility (a11y)", "Prototyping"],
    "product-manager": ["SQL Basics", "A/B Testing", "Figma"],
    "marketing": ["Copywriting", "Google Analytics", "Canva"],
    "hr": ["Excel", "ATS Tools", "Employee Relations"],
    "finance": ["Power BI", "SQL Basics", "Excel"],
    "qa": ["SQL Basics", "CI/CD Basics", "Documentation"],
}

COMPANIES = ["TestCorp", "DemoSoft", "SampleLabs"]
UNIVERSITIES = ["Demo University", "Sample Institute of Technology", "Example State University"]
DEGREES = ["B.Tech (Sample Program)", "B.Sc. (Demo Program)", "BBA (Sample Program)",
           "B.Com (Demo Program)", "B.Des (Sample Program)", "BCA (Demo Program)"]

# Obviously-fake identities: diverse-origin first names + fake last names.
FIRST_NAMES = ["Aarav", "Mei", "Sofia", "Kwame", "Yuki", "Priya", "Omar", "Lena",
               "Diego", "Fatima", "Ravi", "Anika", "Chen", "Nora", "Kofi", "Sana",
               "Rahul", "Isha", "Tunde", "Mira"]
LAST_NAMES = ["Demo", "Testerson", "Sample", "Fakewell", "Mockman", "Exampleton", "Trial", "Placeholder"]

SUMMARIES = [
    "Motivated {title} with hands-on practice-project experience and strong fundamentals.",
    "Detail-oriented {title} seeking to apply {skill} skills on demo and sample projects.",
    "Enthusiastic {title} with coursework plus internship-style practice at sample teams.",
    "Results-focused {title} experienced in {skill} through guided practice projects.",
]

METRIC_BULLETS = [
    "Improved {thing} by {n}% by applying {skill} on a guided practice module",
    "Reduced {thing} time by {n}% after reworking the approach with {skill}",
    "Built a {thing} demo serving {n}+ sample records using {skill}",
    "Raised {thing} coverage/score to {n}% across practice assignments using {skill}",
    "Cut {thing} errors by {n}% by adding checks with {skill}",
]

PLAIN_BULLETS = [
    "Collaborated with a sample team to {task} during a practice sprint",
    "Participated in reviews and {task} with guidance from a mentor",
    "Documented the steps taken to {task} for the team practice wiki",
    "Assisted senior practice-team members to {task} on a demo module",
    "Presented practice findings after helping to {task}",
]

THINGS = ["dashboard load time", "report generation", "test coverage", "data-entry",
          "page render", "query runtime", "onboarding checklist completion", "sample backlog throughput"]


# ---------------------------------------------------------------------------
# Synthetic data builders (seeded RNG only)
# ---------------------------------------------------------------------------

def role_slug(role: str) -> str:
    return role.replace("-", "_")


def build_resume_data(rng: random.Random, idx: int, role: str, senior: bool) -> dict:
    first = rng.choice(FIRST_NAMES)
    last = rng.choice(LAST_NAMES)
    name = f"{first} {last}"
    email = f"{first.lower()}.{last.lower()}{rng.randint(1, 99)}@example.com"
    # Variation axis: some resumes missing phone / links.
    phone = f"+91-90000-{rng.randint(10000, 99999):05d}" if (idx % 4 != 0) else ""
    links = "" if (idx % 3 == 0) else f"portfolio.example.com/{first.lower()}{idx:02d} | github.example.com/{first.lower()}-demo"
    title = ROLE_TITLES[role][1] if senior else ROLE_TITLES[role][0]
    skills_pool = ROLE_SKILLS[role]
    n_skills = rng.randint(4, min(10, len(skills_pool)))
    skills = rng.sample(skills_pool, n_skills)
    summary = rng.choice(SUMMARIES).format(title=ROLE_DISPLAY[role], skill=skills[0])

    n_exp = rng.randint(2, 3)
    experience = []
    for e in range(n_exp):
        company = COMPANIES[(idx + e) % len(COMPANIES)]
        years = f"20{18 + ((idx + e) % 6)} – 20{19 + ((idx + e) % 6)}"
        exp_title = title if e == 0 else rng.choice([ROLE_TITLES[role][0], f"{ROLE_DISPLAY[role]} Intern (Sample)"])
        skill_a, skill_b = rng.sample(skills, min(2, len(skills)))
        task = rng.choice(ROLE_TASKS[role])
        thing = rng.choice(THINGS)
        n = rng.randint(10, 60)
        metric_bullet = rng.choice(METRIC_BULLETS).format(thing=thing, n=n, skill=skill_a)
        plain_bullet = rng.choice(PLAIN_BULLETS).format(task=task)
        mid_bullet = f"Used {skill_b} to support {task} at {company} (practice project)"
        bullets = [metric_bullet, plain_bullet]
        if rng.random() < 0.6:
            bullets.insert(1, mid_bullet)
        experience.append({"title": exp_title, "company": company,
                           "years": years, "bullets": bullets})

    n_proj = rng.randint(1, 2)
    projects = []
    for p in range(n_proj):
        ps = rng.sample(skills, min(3, len(skills)))
        projects.append({
            "name": f"Sample {ROLE_DISPLAY[role]} Project {p + 1}",
            "desc": f"Practice build using {', '.join(ps)}; documented results on demo data.",
        })

    education = {"degree": rng.choice(DEGREES), "school": rng.choice(UNIVERSITIES),
                 "year": f"20{15 + (idx % 9)}"}
    n_certs = rng.choice([0, 1, 1, 2])
    certs = rng.sample(ROLE_CERTS[role], min(n_certs, len(ROLE_CERTS[role])))
    return {"name": name, "email": email, "phone": phone, "links": links,
            "title": title, "summary": summary, "skills": skills,
            "experience": experience, "projects": projects,
            "education": education, "certs": certs, "role": role}


def resume_to_text(d: dict) -> str:
    lines = [d["name"], d["title"]]
    contact = " | ".join(p for p in [d["email"], d["phone"], d["links"]] if p)
    lines.append(contact)
    lines += ["", "SUMMARY", d["summary"], "", "SKILLS", ", ".join(d["skills"]), "", "EXPERIENCE"]
    for e in d["experience"]:
        lines.append(f'{e["title"]} — {e["company"]} ({e["years"]})')
        for b in e["bullets"]:
            lines.append(f"- {b}")
        lines.append("")
    lines.append("PROJECTS")
    for p in d["projects"]:
        lines.append(f'- {p["name"]}: {p["desc"]}')
    lines += ["", "EDUCATION", f'{d["education"]["degree"]}, {d["education"]["school"]} ({d["education"]["year"]})']
    if d["certs"]:
        lines += ["", "CERTIFICATIONS"] + [f"- {c}" for c in d["certs"]]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# PDF renderers (ReportLab Platypus, 3 layout variants)
# ---------------------------------------------------------------------------

def _para_styles(variant: str):
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    ss = getSampleStyleSheet()
    title = ss["Title"]
    title.fontSize = 20 if variant != "compact" else 16
    h = ss["Heading2"]
    h.fontSize = 13
    h.spaceBefore = 10
    h.spaceAfter = 4
    body = ss["BodyText"]
    body.fontSize = 10
    body.leading = 14
    body.alignment = TA_LEFT
    return title, h, body


def render_pdf(d: dict, path: Path, variant: str, multipage: bool = False) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, HRFlowable, PageBreak)
    title_s, h_s, body_s = _para_styles(variant)
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    st: list = []
    contact = " | ".join(p for p in [d["email"], d["phone"], d["links"]] if p)
    if variant == "classic":
        st += [Paragraph(d["name"], title_s), Paragraph(d["title"], body_s),
               Paragraph(contact, body_s), HRFlowable(width="100%", thickness=1),
               Paragraph("Summary", h_s), Paragraph(d["summary"], body_s),
               Paragraph("Skills", h_s), Paragraph(", ".join(d["skills"]), body_s),
               Paragraph("Experience", h_s)]
        for e in d["experience"]:
            st.append(Paragraph(f'<b>{e["title"]}</b> — {e["company"]} ({e["years"]})', body_s))
            for b in e["bullets"]:
                st.append(Paragraph(f"• {b}", body_s))
            st.append(Spacer(1, 4))
    elif variant == "modern":
        header = Table([[Paragraph(f'<b><font size=18>{d["name"]}</font></b><br/>{d["title"]}', body_s),
                         Paragraph(contact.replace(" | ", "<br/>"), body_s)]],
                       colWidths=[3.5 * inch, 3.0 * inch])
        header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                    ("LEFTPADDING", (0, 0), (-1, -1), 6)]))
        st += [header, HRFlowable(width="100%", thickness=1),
               Paragraph("Summary", h_s), Paragraph(d["summary"], body_s),
               Paragraph("Skills", h_s)]
        rows = [d["skills"][i:i + 2] for i in range(0, len(d["skills"]), 2)]
        t = Table([[Paragraph(c, body_s) for c in r] for r in rows],
                  colWidths=[3.25 * inch, 3.25 * inch])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                               ("LEFTPADDING", (0, 0), (-1, -1), 6)]))
        st += [t, Paragraph("Experience", h_s)]
        for e in d["experience"]:
            t2 = Table([[Paragraph(f'<b>{e["title"]}</b>', body_s),
                         Paragraph(f'{e["company"]} ({e["years"]})', body_s)]],
                       colWidths=[3.25 * inch, 3.25 * inch])
            t2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee"))]))
            st.append(t2)
            for b in e["bullets"]:
                st.append(Paragraph(f"• {b}", body_s))
            st.append(Spacer(1, 4))
    else:  # compact two-section
        left = [Paragraph("<b>SKILLS</b>", body_s)] + [Paragraph(f"• {s}", body_s) for s in d["skills"]]
        left += [Paragraph("<b>EDUCATION</b>", body_s),
                 Paragraph(f'{d["education"]["degree"]}<br/>{d["education"]["school"]}<br/>{d["education"]["year"]}', body_s)]
        if d["certs"]:
            left += [Paragraph("<b>CERTS</b>", body_s)] + [Paragraph(f"• {c}", body_s) for c in d["certs"]]
        right = [Paragraph(f"<b><font size=15>{d['name']}</font></b><br/>{d['title']}<br/>{contact}", body_s),
                 Paragraph("<b>SUMMARY</b>", body_s), Paragraph(d["summary"], body_s),
                 Paragraph("<b>EXPERIENCE</b>", body_s)]
        for e in d["experience"]:
            right.append(Paragraph(f'<b>{e["title"]}</b> — {e["company"]} ({e["years"]})', body_s))
            for b in e["bullets"]:
                right.append(Paragraph(f"• {b}", body_s))
        body = Table([[left, right]], colWidths=[2.2 * inch, 4.3 * inch])
        body.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        st.append(body)
    # Shared tail (projects / education / certs) for classic + modern.
    if variant in ("classic", "modern"):
        st += [Paragraph("Projects", h_s)]
        for p in d["projects"]:
            st.append(Paragraph(f'<b>{p["name"]}</b>: {p["desc"]}', body_s))
        st += [Paragraph("Education", h_s),
               Paragraph(f'{d["education"]["degree"]}, {d["education"]["school"]} ({d["education"]["year"]})', body_s)]
        if d["certs"]:
            st += [Paragraph("Certifications", h_s)] + [Paragraph(f"• {c}", body_s) for c in d["certs"]]
    else:
        st += [Paragraph("Projects", h_s)]
        for p in d["projects"]:
            st.append(Paragraph(f'<b>{p["name"]}</b>: {p["desc"]}', body_s))
    if multipage:
        # Variation axis: one multi-page resume — pad with extra practice entries.
        st.append(PageBreak())
        st += [Paragraph("Additional Practice Experience (page 2)", h_s),
               Paragraph("Extra sample entries included so this resume spans two pages.", body_s)]
        for i in range(6):
            st.append(Paragraph(f"<b>Practice Module {i + 1}</b> — SampleLabs (demo rotation)", body_s))
            st.append(Paragraph(f"• Completed guided exercises covering {d['skills'][i % len(d['skills'])]} on synthetic demo data.", body_s))
    doc.build(st)


# ---------------------------------------------------------------------------
# DOCX renderer (python-docx; ~2 files use tables, varied headings)
# ---------------------------------------------------------------------------

HEADING_SETS = [
    {"summary": "Professional Summary", "exp": "Work Experience", "skills": "Technical Skills"},
    {"summary": "Summary", "exp": "Experience", "skills": "Skills"},
    {"summary": "Profile", "exp": "Employment History", "skills": "Core Competencies"},
]


def render_docx(d: dict, path: Path, use_table: bool, heading_set: dict) -> None:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.add_heading(d["name"], level=0)
    doc.add_paragraph(d["title"])
    contact = " | ".join(p for p in [d["email"], d["phone"], d["links"]] if p)
    if contact:
        doc.add_paragraph(contact)
    doc.add_heading(heading_set["summary"], level=1)
    doc.add_paragraph(d["summary"])
    doc.add_heading(heading_set["skills"], level=1)
    doc.add_paragraph(", ".join(d["skills"]))
    doc.add_heading(heading_set["exp"], level=1)
    if use_table:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Role", "Company", "Years"
        for e in d["experience"]:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = e["title"], e["company"], e["years"]
            for b in e["bullets"]:
                doc.add_paragraph(b, style="List Bullet")
    else:
        for e in d["experience"]:
            doc.add_heading(f'{e["title"]} — {e["company"]} ({e["years"]})', level=2)
            for b in e["bullets"]:
                doc.add_paragraph(b, style="List Bullet")
    doc.add_heading("Projects", level=1)
    for p in d["projects"]:
        doc.add_paragraph(f'{p["name"]}: {p["desc"]}', style="List Bullet")
    doc.add_heading("Education", level=1)
    doc.add_paragraph(f'{d["education"]["degree"]}, {d["education"]["school"]} ({d["education"]["year"]})')
    if d["certs"]:
        doc.add_heading("Certifications", level=1)
        for c in d["certs"]:
            doc.add_paragraph(c, style="List Bullet")
    for p in doc.paragraphs:
        for run in p.runs:
            run.font.size = Pt(11)
    doc.save(str(path))


# ---------------------------------------------------------------------------
# "Scanned" PDFs — PIL grayscale render + speckle + rotation, embedded via ReportLab
# ---------------------------------------------------------------------------

def render_scanned(d: dict, path: Path, rng: random.Random) -> None:
    text = resume_to_text(d)
    try:
        from PIL import Image, ImageDraw, ImageFont
        pil_ok = True
    except ImportError:
        pil_ok = False
    if not pil_ok:
        # Fallback: pure-ReportLab "noisy" PDF (jittered text + noise lines).
        from reportlab.pdfgen import canvas as _canvas
        from reportlab.lib.pagesizes import A4 as _A4
        c = _canvas.Canvas(str(path), pagesize=_A4)
        c.setFont("Helvetica", 9)
        y = 800
        jrng = rng
        for line in text.split("\n"):
            x = 40 + jrng.uniform(-3, 3)
            c.drawString(x, y, line[:110])
            y -= 13
            if y < 40:
                c.showPage()
                c.setFont("Helvetica", 9)
                y = 800
        c.setStrokeColorRGB(0.8, 0.8, 0.8)
        for _ in range(40):
            c.line(jrng.uniform(0, 595), jrng.uniform(0, 842),
                   jrng.uniform(0, 595), jrng.uniform(0, 842))
        c.save()
        return
    from PIL import Image, ImageDraw, ImageFont
    try:
        font = ImageFont.load_default(size=26)
        small = ImageFont.load_default(size=22)
    except TypeError:  # older Pillow without size kwarg
        font = small = ImageFont.load_default()
    W, H = 1240, 1754
    img = Image.new("L", (W, H), 255)
    dr = ImageDraw.Draw(img)
    wrapped: list[str] = []
    for para in text.split("\n"):
        wrapped.extend(textwrap.wrap(para, width=72) or [""])
    y = 70
    for i, line in enumerate(wrapped):
        if y > H - 60:
            break
        dr.text((80, y), line, fill=20, font=font if i < 3 else small)
        y += 40
    # Speckle noise (seeded) + slight rotation => scanned look.
    for _ in range(4000):
        x, yy = rng.randint(0, W - 1), rng.randint(0, H - 1)
        dr.point((x, yy), fill=rng.randint(0, 120))
    angle = rng.uniform(-1.2, 1.2)
    img = img.rotate(angle, expand=True, fillcolor=255)
    tmp_png = path.with_suffix(".tmp.png")
    img.save(str(tmp_png))
    from reportlab.pdfgen import canvas
    pw, ph = img.width * 0.48, img.height * 0.48
    c = canvas.Canvas(str(path), pagesize=(pw, ph))
    c.drawImage(str(tmp_png), 0, 0, width=pw, height=ph)
    c.save()
    tmp_png.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# JD builder (.txt, junior + senior mix, partial skill overlap)
# ---------------------------------------------------------------------------

def build_jd(rng: random.Random, j: int, role: str, level: str) -> tuple[str, str]:
    junior_title, senior_title = ROLE_TITLES[role]
    req = rng.sample(ROLE_SKILLS[role], 5 if level == "junior" else 6)
    remaining = [s for s in ROLE_SKILLS[role] if s not in req]
    nice = rng.sample(remaining, 2) + rng.sample(ADJACENT_SKILLS[role], 1)
    exp = "0-2 years (practice/internship/sample work welcome)" if level == "junior" else "3+ years (including sample/demo project leadership)"
    title = junior_title if level == "junior" else senior_title
    body = (
        f"Title: {title} (Sample Posting — fictitious company)\n"
        f"Company: {rng.choice(COMPANIES)} (demo listing, not a real opening)\n"
        f"Level: {level}\n\n"
        f"About: {rng.choice(COMPANIES)} Demo Team is a fictitious sample group. "
        f"This listing exists only for resume-matching tests.\n\n"
        f"Responsibilities:\n"
        + "".join(f"- {rng.choice(ROLE_TASKS[role]).capitalize()} with the demo team\n" for _ in range(5))
        + f"\nRequired skills: {', '.join(req)}\n"
        f"Nice-to-have: {', '.join(nice)}\n"
        f"Experience: {exp}\n"
    )
    fname = f"jd_{j + 1:02d}_{role_slug(role)}_{level}.txt"
    return fname, body


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description="Deterministic synthetic test-dataset generator (no network, no LLM).")
    ap.add_argument("--out", default="test-data")
    ap.add_argument("--resumes", type=int, default=50)
    ap.add_argument("--jds", type=int, default=20)
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    n, m, seed = args.resumes, args.jds, args.seed

    # Format split: ~70% PDF / ~20% DOCX / ~10% scanned (exact for 50 -> 35/10/5).
    n_scanned = int(n * 0.1 + 0.5)
    n_docx = int(n * 0.2 + 0.5)
    n_scanned = min(n_scanned, n)
    n_docx = min(n_docx, n - n_scanned)
    n_pdf = n - n_scanned - n_docx
    formats = ["pdf"] * n_pdf + ["docx"] * n_docx + ["scanned"] * n_scanned
    fmt_rng = random.Random(seed + 999)
    fmt_rng.shuffle(formats)
    variants = ["classic", "modern", "compact"]

    counts_by_format = {"pdf": 0, "docx": 0, "scanned": 0}
    counts_by_role: dict[str, int] = {r: 0 for r in ROLE_FAMILIES}
    table_docx_used = 0

    for i in range(n):
        role = ROLE_FAMILIES[i % len(ROLE_FAMILIES)]
        senior = (i // len(ROLE_FAMILIES)) % 2 == 1  # alternate junior/senior cohorts
        rng = random.Random(seed * 100003 + i * 1013)
        d = build_resume_data(rng, i, role, senior)
        fmt = formats[i]
        slug = role_slug(role)
        if fmt == "pdf":
            variant = variants[fmt_rng.randint(0, 2)]
            multipage = (i == n - 1)  # one multi-page resume
            fname = f"resume_{i + 1:03d}_{slug}.pdf"
            render_pdf(d, out / fname, variant, multipage=multipage)
            counts_by_format["pdf"] += 1
        elif fmt == "docx":
            use_table = table_docx_used < 2 and (i % 5 == 0 or table_docx_used == 0 and i > n - 6)
            if table_docx_used >= 2:
                use_table = False
            if use_table:
                table_docx_used += 1
            hs = HEADING_SETS[i % len(HEADING_SETS)]
            fname = f"resume_{i + 1:03d}_{slug}.docx"
            render_docx(d, out / fname, use_table=use_table, heading_set=hs)
            counts_by_format["docx"] += 1
        else:
            fname = f"resume_{i + 1:03d}_{slug}.pdf"
            render_scanned(d, out / fname, rng)
            counts_by_format["scanned"] += 1
        counts_by_role[role] += 1

    for j in range(m):
        role = ROLE_FAMILIES[j % len(ROLE_FAMILIES)]
        level = "junior" if (j // len(ROLE_FAMILIES) + j) % 2 == 0 else "senior"
        jrng = random.Random(seed * 777 + j * 131)
        fname, body = build_jd(jrng, j, role, level)
        (out / fname).write_text(body, encoding="utf-8")

    manifest = {
        "generated_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "seed": seed,
        "counts": {"resumes": n, "jds": m},
        "counts_by_format": counts_by_format,
        "counts_by_role": counts_by_role,
        "provenance": "100% synthetic templates",
        "safety": "free-tier safe, no real personal data",
    }
    (out / "dataset_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    readme = (
        "# Synthetic Test Dataset (private bucket)\n\n"
        "Provenance: 100% synthetic templates + seeded RNG. No network calls, no LLM calls.\n"
        "All identities are obviously fake (names like Aarav Demo, companies TestCorp / DemoSoft / SampleLabs,\n"
        "emails @example.com, phones +91-90000-XXXXX). Safe for free-tier pipelines.\n\n"
        f"Seed: {seed} | Resumes: {n} | JDs: {m}\n\n"
        "## Format breakdown\n\n"
        f"- PDF (ReportLab Platypus, 3 layout variants): {counts_by_format['pdf']}\n"
        f"- DOCX (python-docx, incl. tables): {counts_by_format['docx']}\n"
        f"- Scanned-style PDF (PIL grayscale + noise + rotation): {counts_by_format['scanned']}\n\n"
        "## Role families (5 resumes each at full size)\n\n"
        + "".join(f"- {r} ({ROLE_DISPLAY[r]})\n" for r in ROLE_FAMILIES)
        + "\n## JDs\n\n20 .txt postings (2 per role family: 1 junior + 1 senior) with required/\n"
        "nice-to-have skills partially overlapping the resume skill lists.\n"
    )
    (out / "README_FOR_BUCKET.md").write_text(readme, encoding="utf-8")

    print(f"\nSynthetic dataset written to: {out} (seed={seed})")
    print("\nBy format:")
    for k, v in counts_by_format.items():
        print(f"  {k:<8} {v}")
    print("\nBy role:")
    for r in ROLE_FAMILIES:
        print(f"  {r:<16} {counts_by_role[r]}")
    print(f"\nJDs: {m} | Manifest + README written.")


if __name__ == "__main__":
    main()
